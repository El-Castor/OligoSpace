from __future__ import annotations

from collections import defaultdict
from collections.abc import Iterable, Mapping
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from oligospace_api.services.sequence_analysis import (
    InvalidSequenceError,
    normalize_sequence,
    reverse_complement,
)

MAIN_SHEET = "Oligo-probes CP database"
DIRECTION_NORMALIZATION = {"F": "Forward", "Forward": "Forward", "Reverse": "Reverse"}


@dataclass(frozen=True)
class ImportIssue:
    severity: str
    code: str
    message: str
    oligo_id: str | None = None
    field: str | None = None
    source_value: object | None = None


@dataclass
class ImportPreview:
    source: str
    candidate_rows: int = 0
    importable_rows: int = 0
    rejected_candidate_rows: int = 0
    allocated_only_rows: int = 0
    skipped_rows: int = 0
    imported_rows: int = 0
    issues: list[ImportIssue] = field(default_factory=list)
    normalized_records: list[dict[str, Any]] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {
            "source": self.source,
            "mode": "preview",
            "candidate_rows": self.candidate_rows,
            "importable_rows": self.importable_rows,
            "rejected_candidate_rows": self.rejected_candidate_rows,
            "allocated_only_rows": self.allocated_only_rows,
            "skipped_rows": self.skipped_rows,
            "imported_rows": self.imported_rows,
            "issues": [asdict(issue) for issue in self.issues],
            "normalized_records": self.normalized_records,
        }


def _is_allocated_only(record: Mapping[str, object | None]) -> bool:
    return bool(record.get("Oligo_ID")) and not any(
        value not in (None, "") for key, value in record.items() if key != "Oligo_ID"
    )


def _entered_length(value: object) -> int:
    if isinstance(value, bool):
        raise TypeError("Boolean is not a valid sequence length")
    if isinstance(value, int | float | str | bytes | bytearray):
        return int(value)
    raise TypeError("Unsupported sequence length value")


def preview_records(
    records: Iterable[Mapping[str, object | None]], *, source: str = "in-memory"
) -> ImportPreview:
    preview = ImportPreview(source=source)
    sequences: defaultdict[str, list[str]] = defaultdict(list)
    locations: defaultdict[str, list[str]] = defaultdict(list)

    for source_row, record in enumerate(records, start=2):
        if _is_allocated_only(record):
            preview.allocated_only_rows += 1
            preview.skipped_rows += 1
            continue
        if not any(value not in (None, "") for value in record.values()):
            continue
        preview.candidate_rows += 1
        oligo_id = str(record.get("Oligo_ID") or "").strip() or None
        name = str(record.get("Oligo_Name") or "").strip()
        source_sequence = record.get("Sequence_5to3")
        has_error = False

        if oligo_id is None or not name:
            preview.issues.append(
                ImportIssue(
                    "error",
                    "missing_mandatory_field",
                    "Oligo_ID and Oligo_Name are required for preview acceptance.",
                    oligo_id,
                )
            )
            has_error = True
        if oligo_id and oligo_id.endswith("EX"):
            preview.issues.append(
                ImportIssue(
                    "error",
                    "template_or_example",
                    "Example/template rows are never imported automatically.",
                    oligo_id,
                )
            )
            has_error = True
        if source_sequence in (None, ""):
            preview.issues.append(
                ImportIssue(
                    "error", "missing_sequence", "Sequence is mandatory.", oligo_id, "Sequence_5to3"
                )
            )
            has_error = True
            normalized = None
        else:
            try:
                normalized = normalize_sequence(str(source_sequence))
            except InvalidSequenceError as exc:
                preview.issues.append(
                    ImportIssue(
                        "error",
                        "invalid_sequence",
                        str(exc),
                        oligo_id,
                        "Sequence_5to3",
                        source_sequence,
                    )
                )
                normalized = None
                has_error = True

        direction_source = str(record.get("Direction") or "").strip()
        direction = DIRECTION_NORMALIZATION.get(direction_source)
        if direction is None and direction_source:
            preview.issues.append(
                ImportIssue(
                    "warning",
                    "ambiguous_direction",
                    "Direction is not in the known source vocabulary.",
                    oligo_id,
                    "Direction",
                    direction_source,
                )
            )
        elif direction_source and direction != direction_source:
            preview.issues.append(
                ImportIssue(
                    "warning",
                    "normalized_direction",
                    f"Direction would be normalized from {direction_source!r} to {direction!r}.",
                    oligo_id,
                    "Direction",
                    direction_source,
                )
            )

        entered_length = record.get("Length_nt")
        if normalized is not None and entered_length not in (None, ""):
            try:
                mismatch = _entered_length(entered_length) != len(normalized)
            except (TypeError, ValueError):
                mismatch = True
            if mismatch:
                message = (
                    f"Entered length {entered_length!r} differs from calculated "
                    f"length {len(normalized)}."
                )
                preview.issues.append(
                    ImportIssue(
                        "warning",
                        "length_mismatch",
                        message,
                        oligo_id,
                        "Length_nt",
                        entered_length,
                    )
                )

        location = str(record.get("Location") or "").strip()
        if location and oligo_id:
            locations[location].append(oligo_id)
        if normalized and oligo_id:
            sequences[normalized].append(oligo_id)

        if has_error:
            preview.rejected_candidate_rows += 1
            preview.skipped_rows += 1
            continue
        preview.importable_rows += 1
        preview.normalized_records.append(
            {
                "source_row": source_row,
                "source_oligo_id": oligo_id,
                "name": name,
                "sequence_original": str(source_sequence),
                "normalized_sequence": normalized,
                "project_source": record.get("Project"),
                "experiment_type_source": record.get("Application"),
                "target_source": record.get("Target"),
                "assay_role": direction,
                "location_source": record.get("Location"),
                "validation_source": record.get("Validation_status"),
            }
        )

    for sequence, ids in sequences.items():
        if len(ids) > 1:
            preview.issues.append(
                ImportIssue(
                    "warning", "exact_sequence_duplicate", f"Exact sequence shared by {ids}."
                )
            )
        reverse = reverse_complement(sequence)
        reverse_ids = sequences.get(reverse, [])
        if reverse != sequence and reverse_ids and sequence < reverse:
            preview.issues.append(
                ImportIssue(
                    "warning",
                    "reverse_complement_duplicate",
                    f"Reverse-complement relationship between {ids} and {reverse_ids}.",
                )
            )
    for location, ids in locations.items():
        if len(ids) > 1:
            message = (
                f"Physical position {location!r} is assigned to {ids}; no location is inferred."
            )
            preview.issues.append(
                ImportIssue(
                    "warning",
                    "storage_conflict",
                    message,
                    source_value=location,
                )
            )
    return preview


def preview_workbook(path: Path) -> ImportPreview:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if MAIN_SHEET not in workbook.sheetnames:
            raise ValueError(f"Required sheet {MAIN_SHEET!r} is missing")
        sheet = workbook[MAIN_SHEET]
        values = sheet.iter_rows(values_only=True)
        headers = [str(value) if value is not None else "" for value in next(values)]
        records = (dict(zip(headers, row, strict=False)) for row in values)
        return preview_records(records, source=path.name)
    finally:
        workbook.close()
